import io
from docx import Document

from app.artifacts.base import AdapterContext
from app.artifacts.docx import DocxAdapter


def test_docx_heading_sections():
    d = Document()
    d.add_heading("هدف", level=1)
    d.add_paragraph("متن هدف")
    d.add_heading("دامنه", level=1)
    d.add_paragraph("متن دامنه")
    buf = io.BytesIO()
    d.save(buf)
    ctx = AdapterContext(
        document_id="1", organization_id="o", wiki_id="w", document_version="1", title="D",
        source_filename="x.docx", language="fa", tags=(), visibility="internal", acl_groups=(),
        artifact_bucket="media", artifact_key="documents/1/Sectioning/ocr/sections.docx",
        artifact_sha256="c" * 64, artifact_etag=None,
    )
    doc = DocxAdapter().parse(buf.getvalue(), ctx)
    assert [s.title for s in doc.sections] == ["هدف", "دامنه"]
