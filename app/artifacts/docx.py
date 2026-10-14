from __future__ import annotations

import io

from docx import Document

from app.artifacts.base import AdapterContext, ArtifactAdapter
from app.artifacts.cleaning import clean_text
from app.models.domain import CanonicalContent, CanonicalDocument


class DocxAdapter(ArtifactAdapter):
    def parse(self, data: bytes, context: AdapterContext) -> CanonicalDocument:
        try:
            doc = Document(io.BytesIO(data))
        except Exception as exc:
            raise ValueError("Artifact is not a readable DOCX file") from exc
        sections: list[CanonicalContent] = []
        title = "Document"
        buffer: list[str] = []

        def flush() -> None:
            nonlocal buffer, title
            text = clean_text("\n\n".join(buffer))
            if text:
                order = len(sections) + 1
                sections.append(
                    CanonicalContent(
                        content_id=f"section-{order:03d}",
                        content_type="section",
                        content_title=clean_text(title) or f"Section {order}",
                        text=text,
                        content_order=order,
                    )
                )
            buffer = []

        for paragraph in doc.paragraphs:
            text = clean_text(paragraph.text)
            if not text:
                continue
            style_name = (paragraph.style.name or "").casefold() if paragraph.style else ""
            if style_name.startswith("heading"):
                flush()
                title = text
            else:
                buffer.append(text)
        flush()
        if not sections:
            raise ValueError("DOCX contains no usable text")
        return CanonicalDocument(
            document_id=context.document_id, organization_id=context.organization_id, wiki_id=context.wiki_id,
            document_version=context.document_version, title=context.title, source_filename=context.source_filename,
            language=context.language, tags=context.tags, visibility=context.visibility,
            artifact_bucket=context.artifact_bucket, artifact_key=context.artifact_key,
            artifact_sha256=context.artifact_sha256, artifact_etag=context.artifact_etag,
            source_kind="legacy_stage2", contents=tuple(sections), source_metadata={},
            stage2_job_id=context.stage2_job_id, stage2_pipeline_version=context.stage2_pipeline_version,
            metadata=context.extra_metadata or {},
        )
